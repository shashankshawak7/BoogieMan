"""
The BoogieMan Engine — Core Backend Logic.
Handles Configuration, Connectivity Probing, LLM Calling, Document Extraction, and PDF Report Export.
"""

import os
import sys
import json
import subprocess
from pathlib import Path
from datetime import datetime
import httpx
from fpdf import FPDF

BASE_DIR = Path(__file__).parent.resolve()
PERSONA_PATH = BASE_DIR / "boogieMan.md"
CONFIG_PATH = BASE_DIR / ".boogie_config.json"
ROASTS_DIR = BASE_DIR / "roasts"
ROASTS_DIR.mkdir(exist_ok=True)

DEFAULT_CATALOG = {
    "LM Studio (Local)": {
        "base_url": "http://127.0.0.1:1234/v1",
        "model": "local-model",
        "api_key": ""
    },
    "Ollama (Local)": {
        "base_url": "http://127.0.0.1:11434/v1",
        "model": "llama3",
        "api_key": ""
    },
    "Google Gemini": {
        "api_key": os.getenv("GEMINI_API_KEY", ""),
        "model": "gemini-2.5-flash"
    },
    "Groq": {
        "api_key": os.getenv("GROQ_API_KEY", ""),
        "base_url": "https://api.groq.com/openai/v1",
        "model": "llama-3.3-70b-versatile"
    },
    "OpenAI": {
        "api_key": os.getenv("OPENAI_API_KEY", ""),
        "base_url": "https://api.openai.com/v1",
        "model": "gpt-4o-mini"
    },
    "Anthropic Claude": {
        "api_key": os.getenv("ANTHROPIC_API_KEY", ""),
        "model": "claude-3-5-sonnet-20241022"
    },
    "Azure OpenAI": {
        "api_key": os.getenv("AZURE_OPENAI_API_KEY", ""),
        "endpoint": "",
        "deployment": "gpt-4o",
        "api_version": "2024-02-15-preview"
    },
    "Custom OpenAI-Compatible": {
        "base_url": "http://127.0.0.1:8000/v1",
        "model": "default",
        "api_key": ""
    }
}


def load_config() -> dict:
    cfg = {
        "active_provider": "LM Studio (Local)",
        "providers": json.loads(json.dumps(DEFAULT_CATALOG)),
        "temperature": 0.7
    }
    if CONFIG_PATH.exists():
        try:
            saved = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
            if "active_provider" in saved and saved["active_provider"]:
                cfg["active_provider"] = saved["active_provider"]
            if "temperature" in saved:
                cfg["temperature"] = saved["temperature"]
            if "providers" in saved and isinstance(saved["providers"], dict):
                for p_name, p_vals in saved["providers"].items():
                    if p_name in cfg["providers"]:
                        cfg["providers"][p_name].update(p_vals)
        except Exception:
            pass
    return cfg


def save_config(cfg: dict):
    CONFIG_PATH.write_text(json.dumps(cfg, indent=2), encoding="utf-8")


def get_persona() -> str:
    if PERSONA_PATH.exists():
        return PERSONA_PATH.read_text(encoding="utf-8")
    return "You are The BoogieMan, a cynical Senior Architect who roasts bad ideas ruthlessly."


def check_connection_live(provider: str, p_cfg: dict) -> tuple[bool, str]:
    if not provider or not p_cfg:
        return False, "Not configured"

    if provider == "LM Studio (Local)":
        base = p_cfg.get("base_url", "http://127.0.0.1:1234/v1").rstrip("/")
        if not base.endswith("/v1"):
            base = f"{base}/v1"
        try:
            with httpx.Client(timeout=1.5) as client:
                res = client.get(f"{base}/models")
                if res.status_code == 200:
                    models = res.json().get("data", [])
                    m_names = [m.get("id") for m in models if "id" in m]
                    m_label = f" ({m_names[0]})" if m_names else ""
                    return True, f"Live & Connected{m_label}"
                return False, f"HTTP {res.status_code}"
        except Exception:
            return False, "Server Offline (Start LM Studio on 127.0.0.1:1234)"

    elif provider == "Ollama (Local)":
        base = p_cfg.get("base_url", "http://127.0.0.1:11434/v1").rstrip("/")
        root_url = base.replace("/v1", "")
        try:
            with httpx.Client(timeout=1.5) as client:
                res = client.get(f"{root_url}/api/tags")
                if res.status_code == 200:
                    models = res.json().get("models", [])
                    m_names = [m.get("name") for m in models if "name" in m]
                    m_label = f" ({m_names[0]})" if m_names else ""
                    return True, f"Live & Connected{m_label}"
                return False, f"HTTP {res.status_code}"
        except Exception:
            return False, "Server Offline (Start Ollama on 127.0.0.1:11434)"

    elif provider == "Google Gemini":
        key = p_cfg.get("api_key", "").strip()
        if not key:
            return False, "API Key Required"
        try:
            with httpx.Client(timeout=2.0) as client:
                res = client.get(f"https://generativelanguage.googleapis.com/v1beta/models?key={key}")
                if res.status_code == 200:
                    return True, "Authenticated"
                return False, f"Auth Error ({res.status_code})"
        except Exception:
            return False, "Network Unreachable"

    elif provider in ["OpenAI", "Groq", "Anthropic Claude", "Azure OpenAI"]:
        key = p_cfg.get("api_key", "").strip()
        if not key:
            return False, "API Key Required"
        return True, "Key Stored"

    elif provider == "Custom OpenAI-Compatible":
        base = p_cfg.get("base_url", "").strip().rstrip("/")
        if not base:
            return False, "Base URL Required"
        return True, "Endpoint Set"

    return False, "Unverified"


def get_git_diff(target_dir=".") -> str:
    try:
        res = subprocess.run(["git", "diff", "HEAD"], cwd=target_dir, capture_output=True, text=True, check=True)
        diff = res.stdout.strip()
        if not diff:
            res = subprocess.run(["git", "diff", "--staged"], cwd=target_dir, capture_output=True, text=True, check=True)
            diff = res.stdout.strip()
        return diff
    except Exception as e:
        return f"Error executing git diff: {e}"


def fetch_remote_diff(url_or_link: str, token: str = "", username: str = "") -> tuple[int, str]:
    target_url = url_or_link.strip()
    if not target_url:
        return 400, "Error: Empty URL provided."

    headers = {
        "User-Agent": "BoogieMan-Architect/1.0",
        "Accept": "application/vnd.github.v3.diff"
    }

    if token.strip():
        if username.strip():
            import base64
            auth_str = base64.b64encode(f"{username.strip()}:{token.strip()}".encode()).decode()
            headers["Authorization"] = f"Basic {auth_str}"
        else:
            headers["Authorization"] = f"Bearer {token.strip()}"
            headers["PRIVATE-TOKEN"] = token.strip()

    if "github.com" in target_url or "gitlab.com" in target_url:
        if not target_url.endswith(".diff") and not target_url.endswith(".patch"):
            target_url = target_url.rstrip("/") + ".diff"

    try:
        with httpx.Client(timeout=30.0, follow_redirects=True) as client:
            resp = client.get(target_url, headers=headers)
            if resp.status_code in [401, 404]:
                return resp.status_code, "Authentication required. Repository is private or credentials missing."
            if resp.status_code != 200:
                return resp.status_code, f"Error fetching diff ({resp.status_code}): {resp.text}"
            diff_text = resp.text.strip()
            if not diff_text:
                return 204, "Diff is empty. No changes detected."
            return 200, diff_text
    except Exception as e:
        return 500, f"Network error fetching diff: {e}"


def extract_text_from_upload(up_file) -> str:
    if up_file is None:
        return ""
    name = up_file.name.lower()
    if name.endswith(".pdf"):
        try:
            import pypdf
            reader = pypdf.PdfReader(up_file)
            pages_text = [page.extract_text() or "" for page in reader.pages]
            return f"### DOCUMENT (PDF): {up_file.name} ({len(pages_text)} pages)\n\n" + "\n\n".join(pages_text)
        except Exception as e:
            return f"Error parsing PDF '{up_file.name}': {e}"
    elif name.endswith(".docx"):
        try:
            import docx
            doc = docx.Document(up_file)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            tables_text = []
            for t in doc.tables:
                for row in t.rows:
                    row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    tables_text.append(" | ".join(row_cells))
            full_text = "\n\n".join(paragraphs)
            if tables_text:
                full_text += "\n\n### TABLES IN DOCUMENT:\n" + "\n".join(tables_text)
            return f"### DOCUMENT (Word .docx): {up_file.name}\n\n{full_text}"
        except Exception as e:
            return f"Error parsing Word file '{up_file.name}': {e}"
    elif name.endswith(".xlsx") or name.endswith(".xls"):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(up_file, data_only=True)
            sheets_text = []
            for sheetname in wb.sheetnames:
                sheet = wb[sheetname]
                rows_text = []
                for row in sheet.iter_rows(values_only=True):
                    row_vals = [str(v).strip() if v is not None else "" for v in row]
                    if any(row_vals):
                        rows_text.append(" | ".join(row_vals))
                if rows_text:
                    sheets_text.append(f"#### Sheet: {sheetname}\n" + "\n".join(rows_text[:300]))
            return f"### SPREADSHEET (Excel): {up_file.name}\n\n" + "\n\n".join(sheets_text)
        except Exception as e:
            return f"Error parsing Excel file '{up_file.name}': {e}"
    else:
        try:
            content = up_file.read().decode("utf-8", errors="replace")
            return f"### DOCUMENT: {up_file.name}\n\n{content}"
        except Exception as e:
            return f"Error reading file '{up_file.name}': {e}"


def call_llm(system_prompt: str, user_prompt: str, provider: str, p_cfg: dict, temp: float = 0.7, chat_history: list = None) -> str:
    messages = [{"role": "system", "content": system_prompt}]
    if chat_history:
        messages.extend(chat_history)
    messages.append({"role": "user", "content": user_prompt})

    # 1. Google Gemini
    if provider == "Google Gemini":
        api_key = p_cfg.get("api_key", "").strip()
        model = p_cfg.get("model", "gemini-2.5-flash").strip()
        if not api_key:
            raise ValueError("Google Gemini API Key is missing. Enter it in the sidebar.")
        url = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent?key={api_key}"
        
        # Format for Gemini REST
        contents = []
        for m in messages:
            if m["role"] == "system":
                continue
            role = "user" if m["role"] == "user" else "model"
            contents.append({"role": role, "parts": [{"text": m["content"]}]})
        
        payload = {
            "system_instruction": {"parts": [{"text": system_prompt}]},
            "contents": contents,
            "generationConfig": {"temperature": temp}
        }
        with httpx.Client(timeout=90.0) as client:
            resp = client.post(url, json=payload)
            if resp.status_code != 200:
                raise RuntimeError(f"Gemini API Error ({resp.status_code}): {resp.text}")
            return resp.json()["candidates"][0]["content"]["parts"][0]["text"]

    # 2. Azure OpenAI
    elif provider == "Azure OpenAI":
        endpoint = p_cfg.get("endpoint", "").strip().rstrip("/")
        api_key = p_cfg.get("api_key", "").strip()
        deployment = p_cfg.get("deployment", "gpt-4o").strip()
        api_ver = p_cfg.get("api_version", "2024-02-15-preview").strip()
        if not endpoint or not api_key:
            raise ValueError("Azure Endpoint and API Key are required.")
        url = f"{endpoint}/openai/deployments/{deployment}/chat/completions?api-version={api_ver}"
        headers = {"api-key": api_key, "Content-Type": "application/json"}
        payload = {"messages": messages, "temperature": temp}
        with httpx.Client(timeout=90.0) as client:
            resp = client.post(url, json=payload, headers=headers)
            if resp.status_code != 200:
                raise RuntimeError(f"Azure Error ({resp.status_code}): {resp.text}")
            return resp.json()["choices"][0]["message"]["content"]

    # 3. Anthropic Claude
    elif provider == "Anthropic Claude":
        api_key = p_cfg.get("api_key", "").strip()
        model = p_cfg.get("model", "claude-3-5-sonnet-20241022").strip()
        if not api_key:
            raise ValueError("Anthropic API key is required.")
        url = "https://api.anthropic.com/v1/messages"
        headers = {"x-api-key": api_key, "anthropic-version": "2023-06-01", "content-type": "application/json"}
        claude_messages = [m for m in messages if m["role"] != "system"]
        payload = {
            "model": model,
            "system": system_prompt,
            "messages": claude_messages,
            "max_tokens": 4096,
            "temperature": temp
        }
        with httpx.Client(timeout=90.0) as client:
            resp = client.post(url, json=payload, headers=headers)
            if resp.status_code != 200:
                raise RuntimeError(f"Anthropic Error ({resp.status_code}): {resp.text}")
            return resp.json()["content"][0]["text"]

    # 4. OpenAI / Groq / Ollama / LM Studio / Custom Local
    else:
        raw_base = p_cfg.get("base_url", "").strip().rstrip("/")
        api_key = p_cfg.get("api_key", "").strip()
        model = p_cfg.get("model", "local-model").strip()

        if provider == "LM Studio (Local)":
            raw_base = raw_base or "http://127.0.0.1:1234/v1"
            api_key = api_key or "lm-studio"
        elif provider == "Ollama (Local)":
            raw_base = raw_base or "http://127.0.0.1:11434/v1"
            api_key = api_key or "ollama"
        elif provider == "Groq":
            raw_base = raw_base or "https://api.groq.com/openai/v1"
        elif provider == "OpenAI":
            raw_base = raw_base or "https://api.openai.com/v1"

        if not raw_base.endswith("/v1"):
            clean_base = f"{raw_base}/v1"
        else:
            clean_base = raw_base

        if not api_key and provider not in ["LM Studio (Local)", "Ollama (Local)"]:
            raise ValueError(f"{provider} requires an API Key.")

        url = f"{clean_base}/chat/completions"
        headers = {"Authorization": f"Bearer {api_key or 'none'}", "Content-Type": "application/json"}
        payload = {"model": model, "messages": messages, "temperature": temp}
        with httpx.Client(timeout=90.0) as client:
            resp = client.post(url, json=payload, headers=headers)
            if resp.status_code != 200:
                raise RuntimeError(f"Endpoint Error ({resp.status_code}) at {url}: {resp.text}")
            return resp.json()["choices"][0]["message"]["content"]


# PDF Report Generation using fpdf2
class BoogieManPDF(FPDF):
    def __init__(self, is_executive: bool = False, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.is_executive = is_executive

    def header(self):
        self.set_font('Helvetica', 'B', 10)
        if self.is_executive:
            self.set_text_color(40, 60, 90)
            self.cell(0, 7, "EXECUTIVE ARCHITECTURAL MEMORANDUM | CONFIDENTIAL", border=False, align="L")
        else:
            self.set_text_color(180, 30, 30)
            self.cell(0, 7, "THE BOOGIEMAN -- OFFICIAL ARCHITECTURAL AUTOPSY REPORT", border=False, align="L")
        self.ln(8)
        self.set_draw_color(180, 180, 180)
        self.line(10, 16, 200, 16)
        self.ln(4)

    def footer(self):
        self.set_y(-15)
        self.set_font('Helvetica', 'I', 8)
        self.set_text_color(130, 130, 130)
        self.cell(0, 10, f"Page {self.page_no()}/{{nb}} -- The BoogieMan Architectural Tribunal", align="C")


def sanitize_for_pdf(text: str) -> str:
    replacements = {
        "—": "-", "–": "-", "“": '"', "”": '"', "‘": "'", "’": "'",
        "•": "*", "…": "...", "👹": "[BOOGIEMAN]", "⚜️": "[EXECUTIVE]",
        "●": "*", "○": "o", "⚠️": "[WARNING]", "⚡": "[ACTIVE]",
        "✅": "[PASS]", "❌": "[FAIL]", "💡": "[IDEA]", "📑": "[RFP]",
        "🗺️": "[ROADMAP]", "📋": "[SPEC]", "📄": "[FILE]", "🔍": "[DIFF]"
    }
    for k, v in replacements.items():
        text = text.replace(k, v)
    return text.encode("latin-1", "replace").decode("latin-1")


def generate_pdf_report(title: str, markdown_content: str, is_executive: bool = False) -> bytes:
    pdf = BoogieManPDF(is_executive=is_executive)
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    # Title Banner
    pdf.set_font("Helvetica", "B", 15)
    if is_executive:
        pdf.set_text_color(25, 45, 75)
    else:
        pdf.set_text_color(180, 100, 20)
    pdf.multi_cell(w=pdf.epw, h=8, text=sanitize_for_pdf(title))
    pdf.ln(2)

    pdf.set_font("Helvetica", "I", 9)
    pdf.set_text_color(100, 100, 100)
    pdf.cell(w=pdf.epw, h=5, text=f"Date: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | Target: Technical Evaluation Report", ln=True)
    pdf.ln(4)

    if is_executive:
        pdf.set_draw_color(40, 70, 110)
    else:
        pdf.set_draw_color(180, 100, 20)
    pdf.set_line_width(0.4)
    pdf.line(10, pdf.get_y(), 200, pdf.get_y())
    pdf.ln(5)

    # Content Parse
    lines = markdown_content.split("\n")
    for raw_line in lines:
        line = sanitize_for_pdf(raw_line.strip())
        if line.startswith("### "):
            pdf.ln(4)
            pdf.set_font("Helvetica", "B", 11)
            if is_executive:
                pdf.set_text_color(30, 65, 110)
            else:
                pdf.set_text_color(160, 90, 20)
            pdf.multi_cell(w=pdf.epw, h=6, text=line.replace("### ", ""))
            pdf.ln(1)
        elif line.startswith("## "):
            pdf.ln(5)
            pdf.set_font("Helvetica", "B", 12)
            pdf.set_text_color(20, 20, 20)
            pdf.multi_cell(w=pdf.epw, h=7, text=line.replace("## ", ""))
            pdf.ln(1)
        elif line.startswith("# "):
            pdf.ln(6)
            pdf.set_font("Helvetica", "B", 14)
            pdf.set_text_color(15, 15, 15)
            pdf.multi_cell(w=pdf.epw, h=8, text=line.replace("# ", ""))
            pdf.ln(2)
        elif line.startswith("* ") or line.startswith("- "):
            pdf.set_font("Helvetica", "", 9.5)
            pdf.set_text_color(40, 40, 40)
            pdf.multi_cell(w=pdf.epw, h=5, text=f"   *  {line[2:].replace('**', '')}")
        elif not line:
            pdf.ln(2)
        else:
            pdf.set_font("Helvetica", "", 9.5)
            pdf.set_text_color(40, 40, 40)
            pdf.multi_cell(w=pdf.epw, h=5, text=line.replace("**", ""))

    return bytes(pdf.output())
